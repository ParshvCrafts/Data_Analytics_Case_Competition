from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parent.parent
FIGURES_DIR = PROJECT_ROOT / "outputs" / "figures"
TABLES_DIR = PROJECT_ROOT / "outputs" / "tables"
DELIVERABLES_DIR = PROJECT_ROOT / "deliverables"
SLIDES_DIR = DELIVERABLES_DIR / "slides"
README_DIR = DELIVERABLES_DIR / "github_readme"
ROOT_README = PROJECT_ROOT / "README.md"
README_COPY = README_DIR / "README.md"
PPTX_PATH = SLIDES_DIR / "nonprofit_ai_readiness_case_competition.pptx"
FINDINGS_PATH = PROJECT_ROOT / "findings.json"


def rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


COLORS = {
    "primary": rgb("#1F4E79"),
    "accent1": rgb("#E8743B"),
    "accent2": rgb("#3CB4AC"),
    "neutral": rgb("#6B6B6B"),
    "background": rgb("#F7F4EE"),
    "black": rgb("#1E1E1E"),
    "white": rgb("#FFFFFF"),
}


def add_background(slide):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["background"]
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_title(slide, text: str, top: float = 0.35, width: float = 11.8, font_size: int = 24):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(width), Inches(0.6))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.bold = True
    run.font.size = Pt(font_size)
    run.font.color.rgb = COLORS["black"]
    return box


def add_text(slide, text: str, left: float, top: float, width: float, height: float, font_size: int = 16, color: str = "black", bold: bool = False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = "Arial"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = COLORS[color]
    return box


def add_bullets(slide, bullets: list[str], left: float, top: float, width: float, height: float, font_size: int = 16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.text = bullet
        p.font.name = "Arial"
        p.font.size = Pt(font_size)
        p.font.color.rgb = COLORS["black"]
    return box


def add_footer(slide, text: str):
    add_text(slide, text, 0.6, 6.85, 11.5, 0.3, font_size=8, color="neutral")


def add_picture(slide, path: Path, left: float, top: float, width: float):
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def add_stat_card(slide, label: str, value: str, left: float, top: float, width: float = 2.2, height: float = 1.1, color: str = "primary"):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["white"]
    shape.line.color.rgb = COLORS[color]
    add_text(slide, value, left + 0.12, top + 0.15, width - 0.24, 0.42, font_size=24, color=color, bold=True)
    add_text(slide, label, left + 0.12, top + 0.62, width - 0.24, 0.25, font_size=10, color="neutral")


def build_deck() -> None:
    findings = json.loads(FINDINGS_PATH.read_text(encoding="utf-8"))
    finding_map = {finding["claim"]: finding["value"] for finding in findings}
    gap_table = pd.read_csv(TABLES_DIR / "want_use_gap_usecase_ranked.csv")
    regional_scorecard = pd.read_csv(TABLES_DIR / "regional_scorecard.csv")
    segments = pd.read_csv(TABLES_DIR / "wug_segment_summary.csv")
    quotes = pd.read_csv(TABLES_DIR / "quotes_stage4_verified.csv")
    persona_summary = pd.read_csv(TABLES_DIR / "personas.csv")
    model_metrics = pd.read_csv(TABLES_DIR / "stats_07_cv_model_metrics.csv")

    global prs
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_text(slide, "Why Nonprofits Want AI but Can't Use It", 0.7, 0.8, 5.7, 1.1, font_size=28, bold=True)
    add_text(slide, "A data infrastructure story from 930 nonprofit survey responses worldwide", 0.7, 1.75, 5.7, 0.5, font_size=18, color="neutral")
    add_text(slide, "Bay Area Data Analytics Case Competition\nGivingTuesday x Gates Foundation track", 0.7, 2.45, 4.2, 0.7, font_size=15, color="primary")
    add_picture(slide, FIGURES_DIR / "fig_06_signature_periodic_table.png", 6.7, 0.9, 5.8)
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 2
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "The headline is simple: nonprofits do not need more AI hype, they need better data plumbing")
    add_stat_card(slide, "want AI help organizing data", "60%", 0.7, 1.2, color="accent1")
    add_stat_card(slide, "already use AI for that", "19%", 3.15, 1.2, color="primary")
    add_stat_card(slide, "Foundation-needed segment", "22%", 5.6, 1.2, color="accent2")
    add_picture(slide, FIGURES_DIR / "fig_06_want_use_gap_diverging.png", 0.75, 2.35, 7.2)
    add_bullets(
        slide,
        [
            "Organizing data has the biggest unmet need in the whole survey, with a 41-point want-use gap.",
            "36.99% of respondents want at least three more AI use cases than they currently use.",
            "209 nonprofits fall into the Foundation-needed segment, meaning demand is real but the basics are not ready yet.",
        ],
        8.25,
        2.45,
        4.3,
        2.8,
        font_size=15,
    )
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 3
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "The sample is global, but it is still a small-org dataset with North America and India doing most of the talking")
    add_stat_card(slide, "survey responses", "930", 0.75, 1.15)
    add_stat_card(slide, "15 staff or fewer", "63%", 3.15, 1.15, color="accent1")
    add_stat_card(slide, "Global North", "62.6%", 5.55, 1.15, color="accent2")
    add_stat_card(slide, "AI Consumers", "56%", 7.95, 1.15, color="primary")
    add_picture(slide, FIGURES_DIR / "fig_00a_continent_distribution.png", 0.8, 2.2, 4.0)
    add_picture(slide, FIGURES_DIR / "fig_00b_org_size_distribution.png", 4.75, 2.2, 3.8)
    add_picture(slide, FIGURES_DIR / "fig_00h_cluster_by_ref.png", 8.7, 2.2, 3.9)
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 4
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "The official report found three personas, and our analysis goes further on where demand and readiness come apart")
    add_text(slide, "What the official report already established", 0.8, 1.2, 5.6, 0.35, font_size=18, color="primary", bold=True)
    add_bullets(
        slide,
        [
            "AI Consumers are the largest cluster (56%), Late Adopters are next (28%), and Skeptics are smallest (15%).",
            "Organizational capacity is a weak readiness proxy.",
            "Those with technical backgrounds tend to be more wary of AI risks.",
            "The biggest unmet AI need is help organizing data.",
        ],
        0.8,
        1.65,
        5.6,
        3.4,
        font_size=15,
    )
    add_text(slide, "Where this deck adds something new", 6.8, 1.2, 5.5, 0.35, font_size=18, color="accent1", bold=True)
    add_bullets(
        slide,
        [
            "It turns the want-use gap into a funder-ready segmentation.",
            "It shows India and Africa are not interchangeable inside the Global South story.",
            "It challenges the tidy 15-staff threshold by showing the steepest staffing jump happens earlier.",
            "It uses open text to show hope and fear living side by side.",
        ],
        6.8,
        1.65,
        5.5,
        3.4,
        font_size=15,
    )
    add_footer(slide, "Sources: GivingTuesday AI Readiness and Adoption in the Nonprofit Sector in 2024; our analysis.")

    # Slide 5
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Late Adopters have the biggest Want-Use Gap, which makes them the most important intervention group")
    add_picture(slide, FIGURES_DIR / "fig_06_wug_distribution.png", 0.8, 1.25, 5.4)
    add_picture(slide, FIGURES_DIR / "fig_06_wug_vs_infra.png", 6.4, 1.25, 5.8)
    add_text(slide, "Late Adopters average a 2.41-point gap, versus 1.81 for Consumers and 0.74 for Skeptics.", 0.9, 6.0, 7.2, 0.4, font_size=15, color="black")
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 6
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Technical respondents are not calmer about AI, they are much more likely to worry about bias")
    add_picture(slide, FIGURES_DIR / "fig_02_risk_by_role_hypothesis.png", 0.8, 1.2, 7.1)
    add_bullets(
        slide,
        [
            "79.5% of Tech/MERL respondents flagged biased AI decisions, versus 53.2% of non-technical peers.",
            "That was the only risk difference that stayed significant after Holm correction (adjusted p = 0.007).",
            "This pushes against the idea that more technical literacy simply reduces fear.",
        ],
        8.15,
        1.6,
        4.1,
        2.6,
        font_size=15,
    )
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 7
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "India and Africa tell different readiness stories even before we compare the Global South to the North")
    add_picture(slide, FIGURES_DIR / "fig_05_regional_scorecard.png", 0.7, 1.2, 12.0)
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis. Treat Latin America and Europe cautiously due to small samples.")

    # Slide 8
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Each extra piece of data infrastructure lifts real-world AI use, even before we get fancy about modeling")
    add_picture(slide, FIGURES_DIR / "fig_08_infrastructure_lift.png", 0.9, 1.35, 6.1)
    add_bullets(
        slide,
        [
            "Organizations with an infrastructure score of 0 average 1.23 AI use cases.",
            "At a score of 5, that rises to 3.49 use cases.",
            "In the controlled logistic model, cloud storage still carries an odds ratio of 1.93 and data agreements 1.58.",
            "Org age and the broad North-South split stop mattering once those operational signals are in the model.",
        ],
        7.35,
        1.55,
        4.5,
        3.4,
        font_size=15,
    )
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 9
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "The 15-staff threshold is not magic, because the first major staffing jump happens earlier")
    add_picture(slide, FIGURES_DIR / "fig_02_staff_threshold.png", 0.85, 1.3, 6.2)
    add_bullets(
        slide,
        [
            "The biggest staffing jump is from 0-5 staff to 6-15 staff, not from 6-15 to 16-30.",
            "That tells us the threshold story is really about when the first technical hire becomes possible.",
            "So '15 staff' is a rough planning heuristic, not a law of nonprofit readiness.",
        ],
        7.45,
        1.7,
        4.3,
        2.8,
        font_size=15,
    )
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 10
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Hope and fear coexist in the open text, and Skeptics sound both more negative and more varied")
    add_picture(slide, FIGURES_DIR / "fig_04_hope_fear_ratio.png", 0.8, 1.25, 6.3)
    add_text(slide, "Verified quotes from respondents", 7.35, 1.35, 4.5, 0.3, font_size=18, color="primary", bold=True)
    quote_lines = []
    for row in quotes.head(4).itertuples():
        quote_lines.append(f'"{row.quote}"')
        quote_lines.append(f"  {row.cluster_label}, {row.region}, {row.role}")
        quote_lines.append("")
    add_text(slide, "\n".join(quote_lines).strip(), 7.35, 1.75, 4.7, 3.9, font_size=12, color="black")
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024 open-text responses, our analysis.")

    # Slide 11
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "The most actionable quadrant is the nonprofits that want more AI but still need their foundation built first")
    add_picture(slide, FIGURES_DIR / "fig_06_wug_segment_matrix.png", 0.85, 1.35, 5.6)
    foundation_row = segments.loc[segments["wug_segment"] == "Foundation-needed"].iloc[0]
    add_bullets(
        slide,
        [
            f"Foundation-needed includes {int(foundation_row['count'])} nonprofits, or {foundation_row['pct']:.2f}% of the sample.",
            "Activation-ready groups can absorb pilots now.",
            "Disengaged groups need AI literacy and low-stakes data practice, not a new product demo.",
            "Already-well-served groups can become peer mentors and testbeds.",
        ],
        6.9,
        1.6,
        4.6,
        3.0,
        font_size=15,
    )
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 12
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Our recommendation is to fund data infrastructure first, tailor tools by segment second, and treat AI readiness as an operating problem")
    add_text(slide, "For funders", 0.9, 1.4, 3.2, 0.3, font_size=18, color="accent1", bold=True)
    add_bullets(slide, ["Shift grant dollars from shiny AI pilots toward cloud storage, policies, and data staffing for Foundation-needed nonprofits."], 0.9, 1.8, 3.4, 1.5, font_size=15)
    add_text(slide, "For tool builders", 4.6, 1.4, 3.2, 0.3, font_size=18, color="primary", bold=True)
    add_bullets(slide, ["Design separate journeys for Activation-ready, Foundation-needed, and Disengaged organizations instead of shipping one generic nonprofit AI offer."], 4.6, 1.8, 3.6, 1.6, font_size=15)
    add_text(slide, "For nonprofits", 8.7, 1.4, 3.2, 0.3, font_size=18, color="accent2", bold=True)
    add_bullets(slide, ["Use the infrastructure score as a self-check before adopting more tools. If the plumbing is weak, fix that before scaling AI."], 8.7, 1.8, 3.7, 1.6, font_size=15)
    add_picture(slide, FIGURES_DIR / "fig_06_signature_periodic_table.png", 0.95, 4.05, 11.25)
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 13
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "The story is useful, but it still comes with sample and method limits we should say out loud")
    add_bullets(
        slide,
        [
            "This is cross-sectional self-report data, so we can rank associations but not claim causation.",
            "North America and India dominate the sample, while Africa, Latin America, Europe, and Australia are much smaller.",
            "The open-text analysis is transparent and reproducible, but the lexicon rules are still simplifications.",
            "Our fresh HDBSCAN rerun only partially recreates the published clusters, which means the original parameter choices matter.",
            "The predictive model is modest (AUC around 0.64), so it is an explanation tool, not a deployment tool.",
        ],
        0.9,
        1.65,
        10.9,
        3.8,
        font_size=16,
    )
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 14
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Appendix: a five-persona lens helps split demand-rich nonprofits from true skeptics")
    add_picture(slide, FIGURES_DIR / "fig_03_persona_parallel.png", 0.85, 1.25, 6.1)
    add_text(slide, persona_summary[["persona_name", "story"]].head(5).to_string(index=False), 7.25, 1.45, 4.7, 4.4, font_size=12, color="black")
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    # Slide 15
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Appendix: the predictive model is modest, which is exactly why we use it for explanation instead of forecasting")
    add_picture(slide, FIGURES_DIR / "fig_07_cv_metrics.png", 0.8, 1.3, 4.0)
    add_picture(slide, FIGURES_DIR / "fig_07_rf_feature_importance.png", 4.95, 1.3, 4.0)
    add_picture(slide, FIGURES_DIR / "fig_07_logit_odds_ratios.png", 9.1, 1.3, 3.8)
    add_text(slide, model_metrics.round(3).to_string(index=False), 0.9, 5.95, 5.5, 0.6, font_size=11, color="black")
    add_footer(slide, "Source: GivingTuesday AI Readiness Survey 2024, n=930, our analysis.")

    SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX_PATH)


def build_readme() -> None:
    gap_table = pd.read_csv(TABLES_DIR / "want_use_gap_usecase_ranked.csv")
    segments = pd.read_csv(TABLES_DIR / "wug_segment_summary.csv")
    regional_scorecard = pd.read_csv(TABLES_DIR / "regional_scorecard.csv")
    quotes = pd.read_csv(TABLES_DIR / "quotes_stage4_verified.csv")

    top_gap = gap_table.iloc[0]
    foundation_row = segments.loc[segments["wug_segment"] == "Foundation-needed"].iloc[0]
    activation_row = segments.loc[segments["wug_segment"] == "Activation-ready"].iloc[0]
    india_row = regional_scorecard.loc[regional_scorecard["region"] == "India"].iloc[0]
    africa_row = regional_scorecard.loc[regional_scorecard["region"] == "Africa"].iloc[0]

    readme_text = f"""# Why Nonprofits Want AI but Can't Use It

This repository contains a Bay Area Data Analytics Case Competition submission built on GivingTuesday's 2024 AI Readiness Survey of 930 nonprofits worldwide. Our main argument is simple: the nonprofit AI bottleneck looks less like a tooling problem and more like a data infrastructure problem, especially for organizations that want more from AI than their operating setup can currently support.

## Thesis

The biggest unmet need in this dataset is not another chatbot. It is help organizing data, plus the staffing, storage, and policy basics that make AI useful instead of frustrating.

## Key Findings

- **Organizing data is the largest unmet AI need.** {top_gap['want_rate'] * 100:.1f}% want it, only {top_gap['use_rate'] * 100:.1f}% currently use it, for a {top_gap['gap'] * 100:.1f}-point gap. See `fig_06_want_use_gap_diverging.png`.
- **Late Adopters carry the biggest Want-Use Gap.** Their average WUG is 2.41, versus 1.81 for AI Consumers and 0.74 for AI Skeptics. See `fig_06_wug_distribution.png`.
- **More than one in five nonprofits sit in the Foundation-needed segment.** {foundation_row['pct']:.2f}% of the sample wants more from AI but still lacks enough infrastructure to support it. See `fig_06_wug_segment_matrix.png`.
- **Technical respondents are more worried about bias, not less.** 79.5% of Tech/MERL respondents flagged biased AI decisions, versus 53.2% of non-technical respondents, and that gap stayed significant after Holm correction. See `fig_02_risk_by_role_hypothesis.png`.
- **India and Africa are not interchangeable inside the Global South story.** India shows a higher mean WUG ({india_row['mean_wug']:.2f}), while Africa shows higher comfort ({africa_row['mean_comfort']:.2f}) and a higher tech-or-MERL staffing rate ({africa_row['pct_has_tech_or_merl']:.1f}%). See `fig_05_regional_scorecard.png`.
- **Infrastructure lifts real AI use.** Organizations with an infrastructure score of 0 average 1.23 AI use cases; at a score of 5, that rises to 3.49. See `fig_08_infrastructure_lift.png`.
- **Hope and fear coexist in the text responses.** One verified quote that captures the caution side: "{quotes.iloc[2]['quote']}". See `fig_04_hope_fear_ratio.png`.

## Headline Figures

![Want-Use Gap](outputs/figures/fig_06_want_use_gap_diverging.png)

![Risk by Role](outputs/figures/fig_02_risk_by_role_hypothesis.png)

![Regional Scorecard](outputs/figures/fig_05_regional_scorecard.png)

![Infrastructure Lift](outputs/figures/fig_08_infrastructure_lift.png)

![Hope vs Fear](outputs/figures/fig_04_hope_fear_ratio.png)

![2x2 Segmentation](outputs/figures/fig_06_wug_segment_matrix.png)

## Deliverables

- Slide deck: `deliverables/slides/nonprofit_ai_readiness_case_competition.pptx`
- Executed notebooks: `notebooks/*_executed.ipynb`
- Figures: `outputs/figures/`
- Tables: `outputs/tables/`
- Stats logs: `outputs/stats/`

## Reproducibility

1. Install dependencies:
   ```bash
   python -m pip install -r requirements.txt
   ```
2. Run the sanity check:
   ```bash
   python scripts/00_sanity_check.py
   ```
3. Review the analysis notebooks in order:
   - `notebooks/01_eda.ipynb`
   - `notebooks/02_hypothesis_tests.ipynb`
   - `notebooks/03_clustering_deep_dive.ipynb`
   - `notebooks/04_text_analysis.ipynb`
   - `notebooks/05_geo_cause_analysis.ipynb`
   - `notebooks/06_want_use_gap_index.ipynb`
   - `notebooks/07_predictive_model.ipynb`
4. Regenerate the polished figure set and deliverables:
   ```bash
   python scripts/stage8_visualization_polish.py
   python scripts/build_stage9_deliverables.py
   ```

## Methodology Notes

- The raw survey file and the clustering file were merged after verifying the row order and cluster labels.
- `cluster3` is the verified 3-group HDBSCAN split used throughout the narrative: AI Consumers, Late Adopters, and AI Skeptics.
- Multi-select risk responses were parsed into explicit binary indicators before testing.
- The open-text pipeline uses transparent preprocessing, VADER sentiment, and a rule-based hope/fear lexicon so every label can be audited.
- The predictive model is deliberately modest and interpretable. Its purpose is to rank drivers of readiness, not to promise high-accuracy forecasting.

## Key Files

- `findings.json` keeps a ledger of verified claims.
- `scripts/analysis_utils.py` is the shared data-loading, plotting, and logging backbone.
- `scripts/stage2_hypothesis_tests.py` through `scripts/stage8_visualization_polish.py` hold the reusable stage logic behind the notebooks.

## Sources

- GivingTuesday, *AI Readiness and Adoption in the Nonprofit Sector in 2024*: <https://ai.givingtuesday.org/ai-readiness-report-2024/>
- GivingTuesday, *AI Readiness Survey Report 2024 India*: <https://ai.givingtuesday.org/ai-readiness-survey-report-2024-india/>
- data.org, *Data Maturity Assessment*: <https://data.org/dma/>
- World Economic Forum, *The AI divide between the Global North and Global South*: <https://www.weforum.org/stories/2023/01/davos23-ai-divide-global-north-global-south/>

## Author

Prepared for the Bay Area Data Analytics Case Competition. Update this section with your team name, names, and contact details before submission.
"""

    README_DIR.mkdir(parents=True, exist_ok=True)
    ROOT_README.write_text(readme_text, encoding="utf-8")
    README_COPY.write_text(
        readme_text.replace("outputs/figures/", "../../outputs/figures/").replace(
            "deliverables/slides/nonprofit_ai_readiness_case_competition.pptx",
            "../slides/nonprofit_ai_readiness_case_competition.pptx",
        ),
        encoding="utf-8",
    )


def main() -> None:
    SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    README_DIR.mkdir(parents=True, exist_ok=True)
    build_deck()
    build_readme()
    print(f"Wrote {PPTX_PATH}")
    print(f"Wrote {ROOT_README}")
    print(f"Wrote {README_COPY}")


if __name__ == "__main__":
    main()
